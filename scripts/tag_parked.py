# -*- coding: utf-8 -*-
"""Tag the test cases that are parked on something outside QA.

Two groups came out of the execution run:

  SQL      88 cases whose condition is a database state no feed file can create.
           Development is providing direct database access on 17 September 2026,
           so these are deferred to a SQL verification pass rather than left open
           as a coverage gap.

  ERRFILE  24 cases whose output is assigned to the error file. Across 84 runs
           the download has never contained one, so access has to be confirmed
           before these can be planned at all.

Neither is a failure of the application and neither is a gap in the test data.
The verdicts stay BLOCKED, which is accurate; this records why and who holds it.
"""
import json
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = r"C:\Users\RSERRANO\PRU ADB"
HERE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(ROOT, "Test Data Scenarios for Demo", "PRU_ADB_E2E_Test_Scenarios_v4.xlsx")
COL = 29                      # first free column

TAGS = {
    "sql": (
        "SQL VERIFICATION",
        "Deferred to direct database verification. The condition is a database state that no feed "
        "file can create, so it cannot be reached through the HR Master Console. Development is "
        "providing database access on 17 September 2026. Verify by setting the state directly, then "
        "confirm the outcome at the assigned destination. Not a coverage gap and not an application "
        "defect.",
        "F3E8D8"),
    "errfile": (
        "PENDING ERRFILE ACCESS",
        "Output is assigned to the error file (errfile-<CCYYMMDD>.csv). Across 84 executed runs the "
        "General Artifacts download has never contained one - it carries CNTLRPT, LOADFILE, "
        "LNAERROR, ADBSKIP, ADBERROR, NEWBD, ADSIMSTR and SQLDUMP only. Development is confirming "
        "whether the file can be accessed. Until then no verdict is possible, whatever the test "
        "data does.",
        "E4E9F2"),
}


def main():
    parked = json.load(open(os.path.join(HERE, "parked.json")))
    lookup = {}
    for key, ids in parked.items():
        for tc in ids:
            lookup[tc] = key

    wb = openpyxl.load_workbook(WB)
    ws = wb["E2E Test Cases"]

    head = ws.cell(3, COL, "Parked - Reason and Owner")
    head.fill = PatternFill("solid", fgColor="0F4C5C")
    head.font = Font(color="FFFFFF", bold=True, size=10)
    head.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[openpyxl.utils.get_column_letter(COL)].width = 64

    tagged = {"sql": 0, "errfile": 0}
    for r in range(4, ws.max_row + 1):
        tc = ws.cell(r, 1).value
        if not tc:
            continue
        key = lookup.get(str(tc).strip())
        if not key:
            continue
        label, text, colour = TAGS[key]
        c = ws.cell(r, COL, "%s - %s" % (label, text))
        c.font = Font(size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = PatternFill("solid", fgColor=colour)
        # make the tag visible on the type column too, where a reader scans first
        ws.cell(r, 5).fill = PatternFill("solid", fgColor=colour)
        tagged[key] += 1

    dv = DataValidation(type="list",
                        formula1='"SQL VERIFICATION,PENDING ERRFILE ACCESS,TWO-PHASE,NOT REACHABLE"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("%s4:%s%d" % (openpyxl.utils.get_column_letter(COL),
                         openpyxl.utils.get_column_letter(COL), ws.max_row))

    # ---- record both on the Coverage Summary
    cs = wb["Coverage Summary"]
    nr = cs.max_row + 1
    rows = [
        ["PARK-01", "Deferred to SQL verification", str(tagged["sql"]),
         "88 test cases whose condition is a database state no feed file can create. Development "
         "is providing database access on 17 September 2026. They are tagged SQL VERIFICATION in "
         "the 'Parked - Reason and Owner' column and are to be verified directly against the "
         "database, then confirmed at the assigned destination. They remain BLOCKED in the verdict "
         "column, which is accurate: no verdict was reachable through the console.",
         "Parked - access 17 Sep"],
        ["PARK-02", "Pending confirmation that ERRFILE can be accessed", str(tagged["errfile"]),
         "24 test cases whose output is assigned to errfile-<CCYYMMDD>.csv. Across 84 executed runs "
         "no such file has ever appeared in the General Artifacts download, which carries eight "
         "artifacts and not that one. Development is checking whether it can be made available. "
         "Until it is, no test data can produce a verdict for these rules.",
         "Blocked on access"],
    ]
    small = Font(size=9)
    wrap = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            c = cs.cell(nr + i, j, v)
            c.font = small
            c.alignment = wrap
        cs.row_dimensions[nr + i].height = 96

    wb.save(WB)
    print("tagged SQL VERIFICATION      :", tagged["sql"])
    print("tagged PENDING ERRFILE ACCESS:", tagged["errfile"])
    print("column %d 'Parked - Reason and Owner' added; PARK-01 and PARK-02 on the Coverage Summary"
          % COL)


if __name__ == "__main__":
    main()
