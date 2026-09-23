# -*- coding: utf-8 -*-
"""How far the BRD's own columns can be trusted as a Java test basis.

Adds a Source Reliability sheet to PRU ADB - Blocked Rules Analysis.xlsx.
Prompted by the confirmation that the modernised process performs no SQL: if a
row describes a SQL condition, it is describing the legacy, not the system under
test.
"""
import json
import os
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = r"C:\Users\RSERRANO\PRU ADB"
HERE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(ROOT, "PRU ADB - Blocked Rules Analysis.xlsx")

SQL = re.compile(r"SQLCODE|SQLSTATE|-811|UTT_[A-Z_]+|DB2|SQL_FAILURE")
BRACKET = re.compile(r"\[([A-Z0-9]{4,6})\]")


def main():
    b = openpyxl.load_workbook(os.path.join(ROOT, "ADB BRD v4 [for QA] (1).xlsx"),
                               data_only=True)["Business Rules"]
    w = openpyxl.load_workbook(os.path.join(ROOT, "ADB BRD v4 RAW  duplicate check - error destination.xlsx"),
                               data_only=True)["Business Rules"]
    J, ROWOF = {}, {}
    for r in range(2, 1001):
        rid = w.cell(r, 2).value
        if rid:
            J[str(rid).strip()] = str(w.cell(r, 24).value or "")
            ROWOF[str(rid).strip()] = r
    COND, CODE = {}, {}
    for r in range(2, 491):
        rid = str(b.cell(r, 2).value or "").strip()
        if rid:
            COND[rid] = " ".join([str(b.cell(r, 8).value or ""), str(b.cell(r, 10).value or "")])
            CODE[rid] = str(b.cell(r, 11).value or "").strip()

    cases = json.load(open(os.path.join(HERE, "..", "data", "test-cases.json")))
    exe = {c["rule"] for c in cases if c["executable"]}
    withtc = {c["rule"] for c in cases}

    cond_sql = {k for k, v in COND.items() if SQL.search(v)}
    java_sql = {k for k, v in J.items() if SQL.search(v)}
    legacy = cond_sql | java_sql

    mismatch = []
    for rid, code in CODE.items():
        m = BRACKET.search(J.get(rid, ""))
        if code and m and m.group(1) != code:
            mismatch.append((rid, code, m.group(1), ROWOF[rid]))
    mismatch.sort(key=lambda x: x[3])

    HDR = PatternFill("solid", fgColor="0F4C5C")
    HDRF = Font(color="FFFFFF", bold=True, size=10)
    SMALL = Font(size=9)
    WRAP = Alignment(wrap_text=True, vertical="top")
    THIN = Side(style="thin", color="D3DDE1")
    BORD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    RED = PatternFill("solid", fgColor="FBE9E7")
    AMB = PatternFill("solid", fgColor="FFF4E5")
    OK = PatternFill("solid", fgColor="E1F0E7")

    wb = openpyxl.load_workbook(WB)
    if "Source Reliability" in wb.sheetnames:
        del wb["Source Reliability"]
    s = wb.create_sheet("Source Reliability", 1)
    s["A1"] = "How far each BRD column can be trusted as a Java test basis"
    s["A1"].font = Font(bold=True, size=15, color="0F4C5C")
    s["A2"] = ("The modernised process performs no SQL. A row that describes a SQL condition is "
               "therefore describing the legacy system, not the system under test.")
    s["A2"].font = Font(italic=True, size=9.5, color="555555")
    s.merge_cells("A2:D2")
    s["A2"].alignment = WRAP

    for i, (h, wd) in enumerate([("Finding", 14), ("What was checked", 44), ("Count", 12),
                                 ("What it means", 104)], 1):
        c = s.cell(4, i, h)
        c.fill = HDR
        c.font = HDRF
        c.border = BORD
        s.column_dimensions[get_column_letter(i)].width = wd

    rows = [
        ("SRC-01", "Rules whose Condition / Trigger cites SQLCODE, a UTT_ table or DB2",
         len(cond_sql),
         "I previously described Condition / Trigger as source-neutral and safe to build tests from. "
         "For these rules it is not: the condition is stated as a legacy database operation. The "
         "business intent may still hold, but the stated condition cannot be taken as the trigger.",
         AMB),
        ("SRC-02", "Rules whose Destination Detail (Java) cites a SQL condition",
         len(java_sql),
         "All of these also name a COBOL program. By contrast, of the rules that proved genuinely "
         "buildable, only one names a COBOL program and none mentions SQL - those name real Java "
         "artefacts (SkippedRecordEntry, ControlReport.forHeaderValidationFailure, "
         "ErrorArtifactExportService). The column is dependable where it describes Java and "
         "unreliable where it describes SQLCODE handling.", RED),
        ("SRC-03", "Rules affected in either column", len(legacy),
         "%d of the 489. Of these, %d have a test case and %d are among the executable cases."
         % (len(legacy), len(legacy & withtc), len(legacy & exe)), RED),
        ("SRC-04", "Executable test cases affected", len(legacy & exe),
         "Checked individually and all five stand. Their Java column reads "
         "'LnaValidationErrorEntry via RecordOutcome.Rejected' with a real LNA code, so the "
         "observable outcome does not depend on SQL. Only the wording of the condition is legacy: "
         "%s." % ", ".join(sorted(legacy & exe)), OK),
        ("SRC-05", "Rules where the code bracketed in the Java column contradicts the BRD's own "
                   "Error / Message Code", len(mismatch),
         "Every one resolves to [B0100], which the reference file defines as ALLSTATE FIRM NAME NOT "
         "VALUED. That cannot be right for rules about designations, date of birth, tax identifier "
         "or producer profile fields. The affected rows sit in one narrow band of the spreadsheet "
         "(rows %d-%d), which is the signature of a fill or copy, not of %d separate decisions. "
         "Treat the bracketed code as unreliable and take the expected code from the BRD's Error / "
         "Message Code column - which is what this suite already does, so no test case is wrong."
         % (mismatch[0][3], mismatch[-1][3], len(mismatch)), RED),
    ]
    r = 5
    for ref, what, n, means, fill in rows:
        for i, v in enumerate([ref, what, n, means], 1):
            c = s.cell(r, i, v)
            c.font = Font(size=10, bold=True) if i == 1 else SMALL
            c.alignment = WRAP
            c.border = BORD
            c.fill = fill
        s.row_dimensions[r].height = 76
        r += 1

    r += 2
    s.cell(r, 1, "The %d contradicting codes" % len(mismatch)).font = Font(bold=True, size=12,
                                                                          color="0F4C5C")
    r += 1
    for i, h in enumerate(["Rule", "BRD Error / Message Code", "Bracketed in the Java column",
                           "Sheet row"], 1):
        c = s.cell(r, i, h)
        c.fill = HDR
        c.font = HDRF
        c.border = BORD
    r += 1
    for rid, code, java, row in mismatch:
        for i, v in enumerate([rid, code, java, row], 1):
            c = s.cell(r, i, v)
            c.font = SMALL
            c.alignment = WRAP
            c.border = BORD
            if i == 3:
                c.fill = RED
        r += 1

    wb.save(WB)
    print("Source Reliability sheet added to", os.path.basename(WB))
    print("   condition cites SQL      :", len(cond_sql))
    print("   Java column cites SQL    :", len(java_sql))
    print("   either                   :", len(legacy), "of 489")
    print("   executable cases affected:", len(legacy & exe), sorted(legacy & exe))
    print("   contradicting codes      :", len(mismatch), "all resolving to [B0100]")


if __name__ == "__main__":
    main()
